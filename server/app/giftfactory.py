"""Gift Factory (C:\\claude projects\\3d\\3d manager): the second PC app next to the
Flag Maker, for runs of 20-500 event gifts. Contract v1 is server/GIFT_FACTORY.md;
Gift Factory's tests/check_crm_link.py holds its side to the same text.

Gift Factory has its own paired key and its own /gf/ routes, so pairing one app
never unpairs the other. Only this server writes requests/{rid}.production (what
Gift Factory reports) and requests/{rid}.guests (the client's guest list from the
portal); firestore.rules refuse both to the app.
"""

import csv
import io
import json
import re
import secrets
import time
from pathlib import Path
from typing import Annotated, Literal
from urllib.parse import quote

import openpyxl
from fastapi import APIRouter, Depends, File, Form, Header, HTTPException, Request, UploadFile
from fastapi.responses import FileResponse, Response
from firebase_admin import firestore
from google.cloud.firestore_v1 import FieldFilter
from openpyxl.styles import Font
from pydantic import BaseModel, ConfigDict, Field, ValidationError, model_validator

from . import firebase, jobs
from .studio import TABS, Upload, User, _customer, _data, _iso, _limit, _sha, create_version, tab_field

router = APIRouter()
LISTED = ("draft", "waiting", "approved", "printing")  # Gift Factory's request list
GUESTS_OPEN = ("draft", "waiting", "approved")  # the client may (re)send a guest list
GUEST_TYPES = {".xlsx": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", ".csv": "text/csv"}
MAX_GUESTS_FILE = 2 * 1024 * 1024
MAX_GUEST_ROWS = 5000
NOTE = "כמות ריקה = עותק אחד · Empty amount = one copy"
NOTE_ROW = re.compile(r"Empty amount = one copy|כמות ריקה = עותק אחד")  # Gift Factory's guests.NOTE


# --- pairing (like the Flag Maker's, in its own file) ------------------------------


def _key_file() -> Path:
    return _data() / "giftfactory_key.json"


@router.post("/studio/giftfactory-key")
def pair_giftfactory(_: User) -> dict:
    """A new key replaces the old one; only its hash is kept."""
    key = secrets.token_urlsafe(24)
    _key_file().write_text(json.dumps({"sha256": _sha(key), "created": time.time()}))
    return {"key": key}


def giftfactory(x_studio_key: Annotated[str, Header()] = "") -> None:
    try:
        saved = json.loads(_key_file().read_text())["sha256"]
    except (OSError, ValueError, KeyError):
        raise HTTPException(401, "Gift Factory is not paired") from None
    if not secrets.compare_digest(saved, _sha(x_studio_key)):
        raise HTTPException(401, "Wrong Gift Factory key")


GiftFactory = Annotated[None, Depends(giftfactory)]


@router.get("/gf/info")
def gf_info(_: GiftFactory) -> dict:
    """Gift Factory sends a proof with views only once this says views: true."""
    return {"contract": 1, "views": True}


def _ref(rid: str):
    """The request, or 404. Ids are Firestore's own (letters and digits), never a path."""
    ref = firebase.db().document(f"requests/{rid}") if rid.isalnum() else None
    snap = ref.get() if ref else None
    if not snap or not snap.exists:
        raise HTTPException(404, "No such request")
    return ref, snap.to_dict() or {}


def _guests(r: dict) -> dict | None:
    g = r.get("guests")
    return {"names": g.get("names"), "gifts": g.get("gifts"), "uploadedAt": _iso(g.get("uploadedAt"))} if g else None


@router.get("/gf/requests")
def gf_requests(_: GiftFactory) -> list[dict]:
    docs = firebase.db().collection("requests").where(filter=FieldFilter("status", "in", list(LISTED))).stream()
    out = [
        {"id": d.id, "title": r.get("title", ""), "customer": r.get("customerName", ""), "qty": r.get("qty", 1),
         "status": r.get("status"), "latestVersion": r.get("latestVersion", 0), "updated": _iso(r.get("updatedAt")),
         "guests": _guests(r)}
        for d in docs if (r := d.to_dict() or {})
    ]
    return sorted(out, key=lambda r: r["updated"] or "", reverse=True)


# --- C2: a version with the name proof sheet -----------------------------------------


class GiftUpload(Upload):
    """The Flag Maker's fields plus names (the name check, a JPEG)."""

    def __init__(
        self,
        preview: Annotated[UploadFile, File()],
        model: Annotated[UploadFile | None, File()] = None,
        recipe: Annotated[UploadFile | None, File()] = None,
        spin: Annotated[UploadFile | None, File()] = None,
        spin_asm: Annotated[UploadFile | None, File()] = None,
        names: Annotated[UploadFile | None, File()] = None,
        note: Annotated[str, Form(max_length=500)] = "",
    ):
        super().__init__(preview, model, recipe, spin, spin_asm, note)
        self.files["names"] = names


class Part(BaseModel):
    # Off contract is refused, not guessed at: a change is contract 2 (GIFT_FACTORY.md).
    model_config = ConfigDict(extra="forbid", strict=True)


Count = Annotated[int, Field(ge=0)]
Day = Annotated[str, Field(pattern=r"^\d{4}-\d{2}-\d{2}$")]


class Views(Part):
    """C2b: cols turns around, a sheet row per pitch (degrees above the gift, row 0 from the top)."""

    cols: int = Field(ge=1)
    pitches: list[Annotated[float, Field(ge=-90, le=90)]] = Field(min_length=1)
    start: tuple[Count, Count]  # [row, col] shown first
    tabs: list[Annotated[str, Field(min_length=1, max_length=100)]] = Field(min_length=1, max_length=20)

    @model_validator(mode="after")
    def fits(self):
        if self.cols * len(self.pitches) > 400:
            raise ValueError("over 400 pictures")
        if self.start[0] >= len(self.pitches) or self.start[1] >= self.cols:
            raise ValueError("start is outside the sheet")
        return self


async def plate_tabs(request: Request) -> tuple[dict | None, dict]:
    """views and the files of tabs 2..20 (preview_k, spin_k): one of each per tab, nothing beyond."""
    form = await request.form()  # already parsed for GiftUpload; Starlette keeps it
    if "views" not in form:
        return None, {}  # the Flag Maker's way, as before
    try:
        views = Views.model_validate_json(form["views"]).model_dump(mode="json")
    except (ValidationError, TypeError) as e:
        raise HTTPException(400, f"views: {e}") from None
    want = {tab_field(kind, k) for k in range(1, len(views["tabs"]) + 1) for kind in ("preview", "spin")}
    for field in sorted(want | (set(form) & set(TABS))):
        got = form.getlist(field)
        if field not in want or len(got) != 1 or not getattr(got[0], "filename", ""):
            raise HTTPException(400, f"{len(views['tabs'])} tabs need one preview and one spin each ({field})")
    return views, {f: form[f] for f in want & set(TABS)}


@router.post("/gf/requests/{rid}/versions")
def upload_from_giftfactory(rid: str, _: GiftFactory, up: Annotated[GiftUpload, Depends()],
                            tabs: Annotated[tuple, Depends(plate_tabs)]) -> dict:
    up.views, extra = tabs
    up.files |= extra
    return create_version(rid, up, "giftfactory")


# --- C1, C4, C6: what Gift Factory reports ------------------------------------------


class Breakdown(Part):
    filament: Count
    power: Count
    wear: Count
    packaging: Count
    labour: Count
    failureMarginPct: Count


class Quote(Part):
    units: Count
    plates: Count
    printHours: float = Field(ge=0)
    finishDate: Day | None
    unitCost: Count
    unitPrice: Count
    total: Count
    breakdown: Breakdown


class Progress(Part):
    platesDone: Count
    platesTotal: Count
    unitsOk: Count
    unitsTarget: Count
    failedUnits: Count
    finishDate: Day | None
    lastEvent: Literal["done", "failed", "queued"]


class Cost(Part):
    filament: Count
    power: Count
    wear: Count
    packaging: Count
    labour: Count
    total: Count


class Production(Part):
    contract: Literal[1]
    job: str = Field(default="", max_length=200)
    quote: Quote | None = None
    progress: Progress | None = None
    actualCost: Cost | None = None


@router.put("/gf/requests/{rid}/production")
def put_production(rid: str, _: GiftFactory, body: Production) -> dict:
    """Each part sent replaces that part; the others stay."""
    ref, r = _ref(rid)
    sent = body.model_dump(exclude_unset=True, exclude={"contract"})
    ref.update({**{f"production.{k}": v for k, v in sent.items()}, "production.updatedAt": firestore.SERVER_TIMESTAMP})
    if p := sent.get("progress"):
        old = (r.get("production") or {}).get("progress") or {}
        failed = p["failedUnits"] - old.get("failedUnits", 0)
        news = [f"Plate {p['platesDone']} of {p['platesTotal']} done"] if p["platesDone"] > old.get("platesDone", 0) else []
        if news or failed > 0:
            news.append(f"{failed} units failed on {r.get('title', '')}" if failed > 0 else r.get("title", ""))
            jobs.send_to(jobs.users(), "production", " · ".join(news), r.get("customerName", ""), {"route": "/request", "id": rid})
    return {"ok": True}


# --- C3: the client's guest list ------------------------------------------------------


def _text(v) -> str:
    if v is None:
        return ""
    if isinstance(v, float) and v.is_integer():
        return str(int(v))
    return str(v).strip()


def _rows(ext: str, data: bytes) -> list[tuple[int, list[str]]]:
    """(row number as the spreadsheet shows it, its cells as text) of every row with text, first sheet."""
    if ext == ".csv":
        try:
            text = data.decode("utf-8-sig")
        except UnicodeDecodeError:
            text = data.decode("cp1255", errors="replace")  # Excel's "CSV" on a Hebrew Windows
        rows = csv.reader(io.StringIO(text))
    else:
        rows = openpyxl.load_workbook(io.BytesIO(data), read_only=True, data_only=True).worksheets[0].iter_rows(values_only=True)
    out = []
    for n, row in enumerate(rows, 1):
        cells = [_text(v) for v in row]
        if any(cells):
            out.append((n, cells))
            if len(out) > MAX_GUEST_ROWS + 1:  # + the header
                raise HTTPException(400, f"Over {MAX_GUEST_ROWS} rows")
    return out


def parse_guests(ext: str, data: bytes) -> dict:
    """The portal template's layout, read the way Gift Factory's app/guests.py reads it:
    header = the first row with text; the name column says name/שם, the amount column
    amount/כמות/qty/quantity (else A and B); empty amount = 1 copy, 0 = not coming."""
    try:
        rows = _rows(ext, data)
    except HTTPException:
        raise
    except Exception:  # a broken xlsx (zip, xml) or csv: whatever the library raises
        raise HTTPException(400, f"The file could not be read as {ext}") from None
    if not rows:
        raise HTTPException(400, "The file is empty")
    (_, header), body = rows[0], rows[1:]
    low = [h.lower() for h in header]

    def column(words: tuple[str, ...], fallback: int) -> int:
        return next((i for i, h in enumerate(low) if any(w in h for w in words)), fallback)

    name_col, amount_col = column(("name", "שם"), 0), column(("amount", "כמות", "qty", "quantity"), 1)
    names = gifts = 0
    skipped = []
    for n, cells in body:
        if NOTE_ROW.search(" ".join(cells)):
            continue
        cell = lambda i: cells[i] if i < len(cells) else ""  # noqa: E731
        if not cell(name_col):
            skipped.append([n, "no name"])
            continue
        raw = cell(amount_col)
        if raw and not re.fullmatch(r"[0-9]+", raw):
            raise HTTPException(400, f'Row {n}: the amount "{raw[:40]}" is not a whole number')
        amount = int(raw) if raw else 1
        if amount == 0:
            skipped.append([n, "amount 0"])
            continue
        names += 1
        gifts += amount
    if not names:
        raise HTTPException(400, "No names in the file")
    return {"names": names, "gifts": gifts, "skipped": skipped}


def guest_template() -> bytes:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Guests"
    ws.sheet_view.rightToLeft = True
    ws.append(["שם / Name", "כמות / Amount"])
    ws.append([NOTE])
    for c in ws[1]:
        c.font = Font(bold=True)
    ws["A2"].font = Font(italic=True, color="808080")
    ws.column_dimensions["A"].width = 40
    ws.column_dimensions["B"].width = 18
    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()


@router.get("/p/{token}/guest-template.xlsx")
def portal_guest_template(token: str) -> Response:
    _customer(token)
    return Response(guest_template(), media_type=GUEST_TYPES[".xlsx"],
                    headers={"Content-Disposition": 'attachment; filename="guests.xlsx"', "Cache-Control": "no-store"})


@router.post("/p/{token}/requests/{rid}/guests")
def portal_guests(token: str, rid: str, file: Annotated[UploadFile, File()]) -> dict:
    """Checked before it is kept: a list that fails here never replaces a good one."""
    cid, c = _customer(token)
    _limit(token)
    try:
        ref, r = _ref(rid)
    except HTTPException:
        r = {}
    if r.get("customerId") != cid:
        raise HTTPException(404, "No such request")
    if r.get("status", "draft") not in GUESTS_OPEN:
        raise HTTPException(409, "The order is already in production")
    ext = Path(file.filename or "").suffix.lower()
    if ext not in GUEST_TYPES:
        raise HTTPException(400, "Save it as .xlsx or .csv")
    data = file.file.read(MAX_GUESTS_FILE + 1)
    if len(data) > MAX_GUESTS_FILE:
        raise HTTPException(413, "The file is over 2 MB")
    summary = parse_guests(ext, data)
    folder = _data() / "guests" / rid
    folder.mkdir(parents=True, exist_ok=True)
    (folder / f"guests{ext}").write_bytes(data)
    for other in GUEST_TYPES.keys() - {ext}:
        (folder / f"guests{other}").unlink(missing_ok=True)
    ref.update({"guests": {"names": summary["names"], "gifts": summary["gifts"],
                           "fileName": Path(file.filename).name[:200], "uploadedAt": firestore.SERVER_TIMESTAMP}})
    jobs.send_to(jobs.users(), "client", f"Guest list from {c.get('name', '')}: {summary['gifts']} gifts",
                 r.get("title", ""), {"route": "/request", "id": rid})
    return summary


def _guest_file(rid: str) -> FileResponse:
    _, r = _ref(rid)
    path = next((p for ext in GUEST_TYPES if (p := _data() / "guests" / rid / f"guests{ext}").is_file()), None)
    if not path:
        raise HTTPException(404, "No guest list")
    name = (r.get("guests") or {}).get("fileName") or path.name
    return FileResponse(path, media_type=GUEST_TYPES[path.suffix], headers={
        "Content-Disposition": f"attachment; filename=\"{path.name}\"; filename*=UTF-8''{quote(name)}"})


@router.get("/gf/requests/{rid}/guests")
def gf_guests(rid: str, _: GiftFactory) -> FileResponse:
    return _guest_file(rid)


@router.get("/studio/requests/{rid}/guests")
def studio_guests(rid: str, _: User) -> FileResponse:
    return _guest_file(rid)
